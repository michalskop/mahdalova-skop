"""Data pro text „Spočítali jsme milion za druhé dítě“ (Demografie, T15).

Zdroj: ČSÚ, Demografická příručka 2024 (https://csu.gov.cz/produkty/demograficka-prirucka-2024, staženo 29. 10. 2026)
- tab. 6.8 Živě narozené děti podle pořadí 1950–2024 → počet druhých dětí v jednotlivých letech
- tab. 6.12 Úhrnná plodnost podle pořadí 1920–2024 → druhé dítě na 100 žen
- ČSÚ, Pohyb obyvatelstva – rok 2025 (předběžné údaje): 77,6 tisíce živě narozených, pořadí za rok 2025 zatím není zveřejněné

Návrh: premiér Andrej Babiš v nedělním videu 13. 9. 2026 – odměna 1 milion korun za druhé dítě, doživotní daňová úleva za třetí
(https://www.denik.cz/cesko/babisovo-reseni-nizke-porodnosti-milion-za-druhe-dite-danove-ulevy-za-treti/).
Podmínky (jednorázově, věk matky, příjem, rozložení do let) nejsou zveřejněné, počítáme proto plošnou výplatu 1 000 000 Kč
za každé živě narozené dítě pořadí 2.

Spuštění: python milion_druhe_dite.py → vypíše hodnoty a uloží milion_druhe_dite.json
"""
import json
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).parent
V = HERE / "vstupy"
ODMENA = 1_000_000  # Kč za jedno narozené dítě pořadí 2


def sheet(name):
    return openpyxl.load_workbook(V / name, data_only=True).active


# --- 6.8: druhé děti podle roku ---
rows = {}
for r in sheet("csu_demograficka_prirucka_2024_tab_6_08.xlsx").iter_rows(values_only=True):
    if r[0] is not None and str(r[0]).strip().isdigit():
        rok = int(str(r[0]).strip())
        rows[rok] = {"celkem": int(r[1]), "prvni": int(r[2]), "druhe": int(r[3]), "treti": int(r[4])}

serie = {rok: rows[rok] for rok in range(2010, 2025)}
d24 = rows[2024]
podil24 = d24["druhe"] / d24["celkem"]
print(f"2024: živě narozených {d24['celkem']:,}, z toho 2. děti {d24['druhe']:,} ({podil24:.1%}), 3. děti {d24['treti']:,}")

# --- cena plošného milionu ---
cena24 = d24["druhe"] * ODMENA
print(f"cena za rok na úrovni 2024: {cena24/1e9:.1f} mld. Kč")

# rok 2025: pořadí zatím nezveřejněno; jen ODHAD při stejném podílu druhých dětí jako v roce 2024
celkem25 = 77_600  # ČSÚ, RI Pohyb obyvatelstva 2025 (zaokrouhleno na stovky, předběžné)
odhad_druhe25 = round(celkem25 * podil24)
print(f"ODHAD 2025: {odhad_druhe25:,} druhých dětí → {odhad_druhe25 * ODMENA / 1e9:.1f} mld. Kč (stejný podíl jako 2024)")

# --- scénáře: kolik by stálo jedno dítě, které by se bez odměny nenarodilo ---
# Dítě „navíc“ = druhé dítě, které by se bez odměny nenarodilo. Milion dostane každé druhé dítě, tedy i ti navíc:
# výplata celkem = (druhé děti dnes + děti navíc) × 1 mil.; cena dítěte navíc = výplata celkem / děti navíc.
def scenar(deti_navic):
    celkem = (d24["druhe"] + deti_navic) * ODMENA
    return {
        "deti_navic": round(deti_navic),
        "navic_pct": round(deti_navic / d24["druhe"] * 100, 1),
        "vyplata_mld": round(celkem / 1e9, 1),
        "cena_na_dite_navic_mil": round(celkem / deti_navic / 1e6, 1),
        # z každých 100 vyplacených milionů míří k dětem navíc:
        "z_100_navic": round(deti_navic / (d24["druhe"] + deti_navic) * 100),
    }

scenare = []
for navic_pct in (1, 5, 10, 20):
    sc = scenar(d24["druhe"] * navic_pct / 100)
    scenare.append(sc)
    print(f"  +{navic_pct} % druhých dětí = {sc['deti_navic']:,} dětí navíc → výplata {sc['vyplata_mld']} mld., {sc['cena_na_dite_navic_mil']} mil. Kč na dítě navíc, ze 100 milionů {sc['z_100_navic']} k dětem navíc")

# kdyby odměna vrátila druhé děti na počet z roku 2021 (rok s nejvyšší plodností od 1991)
navrat = rows[2021]["druhe"] - d24["druhe"]
sc = scenar(navrat)
print(f"návrat na úroveň 2021: +{navrat:,} druhých dětí (+{sc['navic_pct']} %) → výplata {sc['vyplata_mld']} mld., {sc['cena_na_dite_navic_mil']} mil. Kč na dítě navíc")
navrat_2021 = sc

# --- druhé dítě na 100 žen (tab. 6.12) a co by znamenalo +x % ---
ws = sheet("csu_demograficka_prirucka_2024_tab_6_12.xlsx")
tfr2 = {}
for r in ws.iter_rows(values_only=True):
    if isinstance(r[0], int) and r[0] in (1990, 2010, 2021, 2024):
        tfr2[r[0]] = round(r[3], 3)
print("úhrnná plodnost 2. pořadí (druhé dítě na ženu):", tfr2)

# vývoj počtu druhých dětí a podílu na všech narozených
print("rok, druhé děti, podíl druhých dětí:")
for rok, v in serie.items():
    print(f"  {rok}: {v['druhe']:,}  {v['druhe']/v['celkem']:.1%}")

# kolik z celkového poklesu narozených od roku 2021 připadá na druhé děti
pokles_celkem = rows[2021]["celkem"] - rows[2024]["celkem"]
pokles_druhe = rows[2021]["druhe"] - rows[2024]["druhe"]
print(f"2021 → 2024: celkem −{pokles_celkem:,}, z toho druhé děti −{pokles_druhe:,} ({pokles_druhe/pokles_celkem:.0%})")

(HERE / "milion_druhe_dite.json").write_text(
    json.dumps({
        "zdroj": "ČSÚ, Demografická příručka 2024, tab. 6.8 a 6.12; ČSÚ, Pohyb obyvatelstva 2025 (předběžné)",
        "odmena_kc": ODMENA,
        "serie_2010_2024": serie,
        "cena_2024_mld": round(cena24 / 1e9, 1),
        "podil_druhych_2024": round(podil24, 4),
        "odhad_2025": {"druhe": odhad_druhe25, "cena_mld": round(odhad_druhe25 * ODMENA / 1e9, 1), "pozn": "odhad při stejném podílu jako 2024; pořadí za 2025 ČSÚ zatím nezveřejnil"},
        "scenare": scenare,
        "navrat_2021": navrat_2021,
        "tfr_druhe": tfr2,
        "pokles_2021_2024": {"celkem": pokles_celkem, "druhe": pokles_druhe},
    }, ensure_ascii=False, indent=1),
    encoding="utf-8",
)
