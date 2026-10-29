"""Data pro článek 05 (načasování rodičovství).

Zdroje (ČSÚ, Demografická příručka 2024, https://csu.gov.cz/produkty/demograficka-prirucka-2024,
staženo 28. 10. 2026):
- tab. 6.11 Míry plodnosti žen podle věku 1950–2024 → kopec plodnosti 1990/2010/2024
- tab. 6.12 Úhrnná plodnost podle pořadí 1920–2024 → první a druhé dítě na 100 žen
- tab. 6.13 Průměrný věk žen při narození dítěte 1950–2024 → věk u 1. a 2. dítěte

Spuštění: python cas_rodicovstvi.py → vypíše hodnoty a uloží cas_rodicovstvi.json
"""
import json
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
V = HERE / "vstupy"
YEARS = (1990, 2010, 2024)


def sheet(name):
    return openpyxl.load_workbook(V / name, data_only=True).active


# 6.11: řádky = věk, sloupce = roky (hlavička v řádku 4)
ws = sheet("csu_demograficka_prirucka_2024_tab_6_11.xlsx")
rows = list(ws.iter_rows(values_only=True))
header = rows[3]
col = {y: header.index(y) for y in YEARS}
kopec = []
for r in rows[4:]:
    age = str(r[0]).strip()
    digits = "".join(ch for ch in age if ch.isdigit())[:2]
    if not digits:
        continue
    vek = int(digits)
    if not 15 <= vek <= 49:
        continue
    for y in YEARS:
        kopec.append({"vek": vek, "rok": y, "na1000": round(float(r[col[y]]), 1)})
for y in YEARS:
    peak = max((k for k in kopec if k["rok"] == y), key=lambda k: k["na1000"])
    print(f"vrchol kopce {y}: věk {peak['vek']}, {peak['na1000']} dětí na 1 000 žen")

# 6.12: úhrnná plodnost podle pořadí → děti daného pořadí na 100 žen
poradi = {}
for r in sheet("csu_demograficka_prirucka_2024_tab_6_12.xlsx").iter_rows(values_only=True):
    if r[0] in (1990, 2024):
        poradi[r[0]] = {"prvni": round(r[2] * 100), "druhe": round(r[3] * 100)}
print("na 100 žen (první, druhé dítě):", poradi)

# 6.13: průměrný věk při 1. a 2. dítěti
vek = {}
for r in sheet("csu_demograficka_prirucka_2024_tab_6_13.xlsx").iter_rows(values_only=True):
    if r[0] in (1990, 2024):
        vek[r[0]] = {"prvni": round(r[2], 1), "druhe": round(r[3], 1)}
print("průměrný věk (první, druhé dítě):", vek)

(HERE / "cas_rodicovstvi.json").write_text(
    json.dumps({"kopec": kopec, "na100zen": poradi, "vek": vek}, ensure_ascii=False, indent=1),
    encoding="utf-8",
)
