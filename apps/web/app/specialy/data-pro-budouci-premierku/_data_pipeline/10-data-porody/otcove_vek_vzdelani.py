"""Muži a načasování dětí: co o otcích říkají oficiální tabulky ČSÚ (podnět Kateřiny, 28. 10. 2026).

Spuštění: python -X utf8 -I otcove_vek_vzdelani.py <300002161116.xlsx> <300002161117.xlsx>
Vstupy (archiv Zdrojová data/CSU/zaostreno-zeny-muzi-2016, viz ZDROJE.md):
  1) ČSÚ, Zaostřeno na ženy a muže 2016, tab. 1-16 „Živě narození podle věku otce i matky ve vybraných letech“ (2007, 2010, 2015)
  2) ČSÚ, tab. 1-17 „Živě narození podle vzdělání matky, otce“ (1993–2015)
Výstupy: data/otcove_vek.csv, data/otcove_vzdelani.csv, data/otcove_vypis.txt

Co skript NEUMÍ: počet dětí mužů podle jejich vzdělání. Tabulky jsou o narozených dětech (otec uvedený na hlášení
o narození), ne o mužích; ve sčítání 2021 se počet dětí zjišťoval jen u žen. Nelze tedy ověřit, zda mají muži s vysokou školou
stejný konečný počet dětí jako ostatní; lze popsat jen věk otců a vzdělání otců novorozených dětí.
Průměrný věk z věkových skupin: středy skupin (do 19 = 17,5; 45+ = 47 u matek; 60+ = 62 u otců), děti s nezjištěným věkem se vynechají.
"""
import csv
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
os.makedirs(DATA, exist_ok=True)
f_vek, f_vzd = sys.argv[1], sys.argv[2]
L = []


def num(x):
    return 0 if x in (None, "-", "") else float(x)


# --- 1) věk otce × věk matky (tab. 1-16)
ws = openpyxl.load_workbook(f_vek, data_only=True).worksheets[0]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
blocks, year = {}, None
for r in rows:
    if isinstance(r[0], int) and all(c is None for c in r[1:]):
        year = r[0]
        blocks[year] = []
    elif year and r[0] and not str(r[0]).startswith(("Věk otce", "Pramen")):
        blocks[year].append(r)
FATHER_MID = {"do 19": 17.5, "20-24": 22, "25-29": 27, "30-34": 32, "35-39": 37, "40-44": 42, "45-49": 47, "50-54": 52, "55-59": 57, "60+": 62}
MOTHER_MID = {"do 19": 17.5, "20-24": 22, "25-29": 27, "30-34": 32, "35-39": 37, "40-44": 42, "45+": 47}
mother_cols = ["do 19", "20-24", "25-29", "30-34", "35-39", "40-44", "45+"]
out = []
for y, bl in blocks.items():
    tot = {r[0].split("/")[0].strip(): r for r in bl}
    fath = {k: num(v[1]) for k, v in tot.items() if k in FATHER_MID}
    celkem = num(tot["Celkem"][1])
    nezj_otec = num(tot["Nezjištěno"][1])
    mean_f = sum(FATHER_MID[k] * n for k, n in fath.items()) / sum(fath.values())
    moth = {k: num(tot["Celkem"][2 + i]) for i, k in enumerate(mother_cols)}
    mean_m = sum(MOTHER_MID[k] * n for k, n in moth.items()) / sum(moth.values())
    share35 = 100 * sum(n for k, n in fath.items() if k in ("35-39", "40-44", "45-49", "50-54", "55-59", "60+")) / sum(fath.values())
    share_unknown = 100 * nezj_otec / celkem
    out.append([y, int(celkem), int(nezj_otec), f"{share_unknown:.1f}", f"{mean_f:.1f}", f"{mean_m:.1f}", f"{share35:.1f}"])
    L.append(f"{y}: narozených {int(celkem):,}, věk otce nezjištěn u {share_unknown:.1f} %; průměrný věk otce (ze skupin) {mean_f:.1f}, "
             f"matky {mean_m:.1f}; podíl otců 35+ let {share35:.1f} %".replace(",", " "))
with open(os.path.join(DATA, "otcove_vek.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["rok", "narozeni", "vek_otce_nezjisten", "vek_otce_nezjisten_pct", "prumerny_vek_otce", "prumerny_vek_matky", "podil_otcu_35plus_pct"])
    w.writerows(out)

# --- 2) vzdělání otce a matky (tab. 1-17)
ws = openpyxl.load_workbook(f_vzd, data_only=True).worksheets[0]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
blocks, year = {}, None
for r in rows:
    if r[0] is None and isinstance(r[1], int) and all(c is None for c in r[2:]):
        year = r[1]
        blocks[year] = []
    elif year and r[0] and "/" in str(r[0]):
        blocks[year].append(r)
LAB = {"celkem": "Celkem", "zakladni": "základní", "bez": "střední bez", "mat": "střední s", "vs": "vysokoškolské", "nezj": "nezjištěno"}
out = []
for y, bl in sorted(blocks.items()):
    d = {}
    for r in bl:
        name = str(r[0]).split("/")[0].strip()
        for k, lab in LAB.items():
            if name.startswith(lab):
                d[k] = r
    tot = num(d["celkem"][1])
    fa_vs, fa_nez = num(d["vs"][1]), num(d["nezj"][1])
    ma_vs = num(d["celkem"][5])
    ma_nez = num(d["celkem"][6])
    p_fa = 100 * fa_vs / (tot - fa_nez)
    p_ma = 100 * ma_vs / (tot - ma_nez)
    out.append([y, int(tot), f"{100 * fa_nez / tot:.1f}", f"{p_fa:.1f}", f"{100 * ma_nez / tot:.1f}", f"{p_ma:.1f}"])
with open(os.path.join(DATA, "otcove_vzdelani.csv"), "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["rok", "narozeni", "otec_nezjisteno_pct", "otec_vs_z_uvedenych_pct", "matka_nezjisteno_pct", "matka_vs_z_uvedenych_pct"])
    w.writerows(out)
for r in out:
    if r[0] in (1993, 2000, 2007, 2010, 2015) or r is out[-1]:
        L.append(f"{r[0]}: narozených {r[1]}, otec s vysokou školou {r[3]} % z uvedených (vzdělání otce nezjištěno u {r[2]} %); "
                 f"matka s vysokou školou {r[5]} % z uvedených (nezjištěno {r[4]} %)")

open(os.path.join(DATA, "otcove_vypis.txt"), "w", encoding="utf-8").write("\n".join(L) + "\n")
print("\n".join(L))
