# Lidskoprávní a sociální profil kandidátů do Senátu 2026 z odpovědí v naší Volební kalkulačce (KohoVolit.eu / DataTimes).
# Vstupy: zdroj/kalkulacka-senat-2026-odpovedi.xlsx (export Google Formu; obsahuje PIN – NIKDY necommitovat ani nevypisovat),
#         zdroj/csv/serk.csv (registr kandidátů ČSÚ 2026: https://volby.gov.cz/opendata/se2026/SE2026reg20261017_csv.zip),
#         senat-istanbul-2024-01-24.txt (hlasování Senátu o Istanbulské úmluvě, opsáno ze senat.cz),
#         číselník stran: nejdřív spustit senat_hist.py v článku volby-cr-2026-10-17-senat-2026-co-ukazaly-volby.
# Index = podíl odpovědí „Pro“ ze zodpovězených otázek v bloku (Nevím/neodpověděl se nepočítá do jmenovatele).
import openpyxl, csv, os, sys, re, unicodedata, collections, statistics
sys.stdout.reconfigure(encoding='utf-8')
H = os.path.dirname(os.path.abspath(__file__))
PRAVA = {8: 'menstruační potřeby ve snížené DPH', 16: 'ratifikovat Istanbulskou úmluvu', 23: 'manželství pro páry stejného pohlaví', 24: 'právo na interrupci v Ústavě'}
SOC = {2: 'stát víc pomáhá s obecními byty', 6: 'vyšší daně pro lidi s vyššími příjmy', 14: 'kurýři a řidiči aplikací jako zaměstnanci', 15: 'minimální mzda aspoň polovina průměrné', 17: 'automaticky zastavovat bezvýsledné exekuce'}
def norm(s): return ''.join(c for c in unicodedata.normalize('NFD', s.lower()) if unicodedata.category(c) != 'Mn')
TIT = re.compile(r'\b(ing|mgr|mudr|phdr|judr|rndr|bc|doc|prof|ph\.?d|mba|dis|csc|dr|mpa|ll\.?m|paeddr|mvdr|thdr|phmr|dipl)\b\.?', re.I)
def klic(jm):
    t = [x for x in re.split(r'[\s,]+', TIT.sub(' ', jm)) if x and len(x) > 1]
    return frozenset(norm(x) for x in t)
# registr ČSÚ
cvs = {}
reg = list(csv.DictReader(open(os.path.join(H, 'zdroj', 'csv', 'serk.csv'), encoding='cp1250'), delimiter=';'))
K = {}
for r in reg:
    k = (int(r['OBVOD']), frozenset([norm(r['JMENO'].split()[0]), *[norm(x) for x in r['PRIJMENI'].split()]]))
    K[k] = r
# odpovědi
ws = openpyxl.load_workbook(os.path.join(H, 'zdroj', 'kalkulacka-senat-2026-odpovedi.xlsx'), read_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
qcol = {int(m.group(1)): i for i, h in enumerate(hdr) if h and (m := re.match(r'(\d+) / 40', str(h)))}
odp = {}
nesparovano = []
for r in rows[1:]:
    jm = str(r[1] or ''); 
    if 'Test' in jm: continue
    kj = klic(jm)
    kand = [v for (o, kk), v in K.items() if kk <= kj or (len(kk & kj) >= 2)]
    if len(kand) != 1:
        nesparovano.append((jm, str(r[2]), len(kand))); continue
    c = kand[0]; key = (c['OBVOD'], c['PRIJMENI'], c['JMENO'])
    odp[key] = (c, {q: (str(r[i]).strip() if r[i] else '') for q, i in qcol.items()}, str(r[0]))  # pozdější řádek přepíše dřívější
print('odpovědí bez testu:', len(rows) - 2, '| spárováno kandidátů:', len(odp), '| nespárováno:', nesparovano)
def idx(a, qs):
    v = [a[q] for q in qs if a[q] in ('Pro', 'Proti')]
    return (sum(x == 'Pro' for x in v) / len(v)) if v else None
cv = {}
for p in ['/'.join([H, 'zdroj', 'csv', 'serk.csv'])]: pass
zvol = lambda c: c['ZVOLEN_K1'] == '1' or c['ZVOLEN_K2'] == '1'
fin = lambda c: zvol(c) or float(c['PROC_K2'] or 0) > 0
print('\nZvolení, kteří kalkulačku vyplnili:', sum(zvol(c) for c, _, _ in odp.values()), 'z 27')
print('Finalisté (2. kolo nebo zvolen v 1. kole), kteří vyplnili:', sum(fin(c) for c, _, _ in odp.values()), 'z', sum(fin(c) for c in reg))
def skup(nazev, f):
    x = [(c, a) for c, a, _ in odp.values() if f(c)]
    p = [idx(a, PRAVA) for c, a in x]; s = [idx(a, SOC) for c, a in x]
    p = [v for v in p if v is not None]; s = [v for v in s if v is not None]
    print(f"{nazev:42s} n={len(x):3d} | práva: průměr {100*statistics.mean(p):5.1f} % | sociální: průměr {100*statistics.mean(s):5.1f} %" if p and s else f'{nazev}: n={len(x)}')
print()
skup('všichni, kdo vyplnili', lambda c: True)
skup('zvolení', zvol)
skup('nezvolení', lambda c: not zvol(c))
skup('poražení ve 2. kole', lambda c: fin(c) and not zvol(c))
print('\n=== jednotlivé otázky: podíl Pro mezi zvolenými × poraženými ve 2. kole × všemi ===')
for q, t in {**PRAVA, **SOC}.items():
    def pr(f):
        v = [a[q] for c, a, _ in odp.values() if f(c) and a[q] in ('Pro', 'Proti')]
        return f"{sum(x == 'Pro' for x in v)}/{len(v)}"
    print(f"{q:2d} {t:45s} zvolení {pr(zvol):7s} poražení 2. k. {pr(lambda c: fin(c) and not zvol(c)):7s} všichni {pr(lambda c: True)}")
print('\n=== zvolení jednotlivě (odpovědi na 8, 16, 23, 24) ===')
for c, a, ts in sorted(odp.values(), key=lambda t: int(t[0]['OBVOD'])):
    if zvol(c) or c['PRIJMENI'] in ('Sucharda Šípová', 'Šípová'):
        print(c['OBVOD'], c['JMENO'], c['PRIJMENI'], '|', ' | '.join(f"{q}:{a[q]}" for q in PRAVA), '| práva', idx(a, PRAVA), '| soc', idx(a, SOC))
print('\nZvolení, kteří nevyplnili:', [f"{c['OBVOD']} {c['PRIJMENI']}" for c in reg if zvol(c) and (c['OBVOD'], c['PRIJMENI'], c['JMENO']) not in odp])

print('\n=== podle pohlaví a strany (navrhující strana ČSÚ) ===')
cvs_p = os.path.join(H, '..', '..', 'volby-cr-2026-10-17-senat-2026-co-ukazaly-volby', '_data_pipeline', 'zdroj', 'dl', 'SENATciselniky20250118_csv', 'csv', 'cvs.csv')
cv = {r['VSTRANA']: r['ZKRATKAV8'] for r in csv.DictReader(open(os.path.normpath(cvs_p), encoding='cp1250'), delimiter=';')}
def zena(c):
    p = c['PRIJMENI'].split()[-1]; j = c['JMENO'].split()[0]
    return p.endswith(('á', 'ová')) or (j.endswith(('a', 'e')) and j not in ('Jiří', 'Ota', 'Nikola'))
skup('ženy', zena); skup('muži', lambda c: not zena(c))
for st in ['ANO', 'SPD', 'ODS', 'STAN', 'KDU-ČSL', 'TOP 09']:
    skup(f'navrhla {st}', lambda c, st=st: cv.get(c['NSTRANA'], c['NSTRANA']) == st)
skup('ODS+STAN+KDU+TOP 09', lambda c: cv.get(c['NSTRANA'], c['NSTRANA']) in ('ODS', 'STAN', 'KDU-ČSL', 'TOP 09'))

# --- Obvod po obvodu: jak k Istanbulské úmluvě hlasoval odcházející senátor (24. 1. 2024) × odpověď nově zvoleného na otázku 16 ---
import json
VOTE = {}
for line in open(os.path.join(H, 'senat-istanbul-2024-01-24.txt'), encoding='utf-8'):
    if line.startswith('#'): continue
    for it in line.strip().split('|'):
        if it: VOTE[it[2:].strip()] = it[0]
VYKLAD = {'A': 'pro', 'N': 'proti', 'X': 'zdržel se', '0': 'nepřítomen'}
sd = json.load(open(os.path.normpath(os.path.join(H, '..', '..', 'volby-cr-2026-10-17-senat-druhe-kolo-predikce', '_data_pipeline', 'senat-data.json')), encoding='utf-8'))
puvodni = {s['obvod']: s['jmeno'] for s in sd['senatori'] if s['konci2026']}
novy = {int(c['OBVOD']): c for c in reg if zvol(c)}
ans = {int(c['OBVOD']): a for c, a, _ in odp.values() if zvol(c)}
print('\n=== 27 obvodů: odcházející senátor – hlas k Istanbulu 2024 | nově zvolený – otázka 16 v kalkulačce ===')
cnt = collections.Counter()
for o in range(3, 82, 3):
    p = puvodni[o]; v = VYKLAD.get(VOTE.get(p, '?'), 'nebyl v Senátu / nenalezen')
    n = novy[o]; a = ans.get(o, {}).get(16) or 'nevyplnil'
    stejny = p.split()[-1] == n['PRIJMENI'].split()[-1]
    cnt[(v, a)] += 1
    print(f"{o:2d} {p:28s} {v:12s} -> {n['JMENO']} {n['PRIJMENI']:24s} {a}{'  (obhájil)' if stejny else ''}")
print('\nodcházející: pro', sum(VOTE.get(p) == 'A' for p in puvodni.values()), '| proti', sum(VOTE.get(p) == 'N' for p in puvodni.values()),
      '| zdržel se', sum(VOTE.get(p) == 'X' for p in puvodni.values()), '| nepřítomen', sum(VOTE.get(p) == '0' for p in puvodni.values()),
      '| nebyl', sum(p not in VOTE for p in puvodni.values()))
print('noví (kalkulačka ot. 16): pro', sum(a.get(16) == 'Pro' for a in ans.values()), '| proti', sum(a.get(16) == 'Proti' for a in ans.values()),
      '| nevím', sum((a.get(16) or '').startswith('Nevím') for a in ans.values()), '| nevyplnili', 27 - len(ans))

# --- Export pro kontrolu: kalkulacka-index-detail.xlsx (bez PIN a bez komentářů kandidátů) ---
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook()
m = wb.active; m.title = 'Metoda'
for r in [
    ['Jak měříme lidskoprávní a sociální profil kandidátů (Volební kalkulačka 2026 – Senát)'],
    [],
    ['Zdroj', 'Odpovědi kandidátů v naší Volební kalkulačce (dotazník, 26. 9.–10. 2026), spárované s registrem kandidátů ČSÚ 2026 podle jména a obvodu.'],
    ['Blok „lidská práva“', 'Otázky 8, 16, 23, 24: ' + '; '.join(f'{q} = {t}' for q, t in PRAVA.items())],
    ['Blok „sociální“', 'Otázky 2, 6, 14, 15, 17: ' + '; '.join(f'{q} = {t}' for q, t in SOC.items())],
    ['Index', 'Podíl odpovědí „Pro“ ze všech odpovědí „Pro“ a „Proti“ v bloku. „Nevím / nemám jednoznačný názor“ a nevyplněné se nepočítají do jmenovatele.'],
    ['Příklad', 'Petr Paták: 8 Pro, 16 Nevím, 23 Nevím, 24 Proti → jmenovatel 2 (Pro + Proti), Pro 1 → index 1/2 = 50 %. Adéla Sucharda Šípová: 8 Pro, 16 Pro, 23 Pro, 24 Nevím → 3/3 = 100 %.'],
    ['Pozor', 'Všechny 4 otázky jsou formulované tak, že „Pro“ = lidskoprávně/genderově progresivní postoj; u sociálního bloku „Pro“ = sociálně citlivější postoj. Index měří deklarované postoje v dotazníku, ne hlasování.'],
    ['Pozor 2', '„Nevím“ nesnižuje index, jen ubírá otázky; kandidát s jedinou odpovědí „Pro“ má 100 %. Proto ve listu Kandidáti uvádíme i počet odpovědí, ze kterých index vznikl.'],
    ['Odcházející senátoři', 'Nemáme jejich kalkulačku, proto pro ně bereme skutečné hlasování o Istanbulské úmluvě v Senátu 24. 1. 2024 (senat.cz, 21. schůze, hlasování 24). Legenda senat.cz: A ano, N ne, X zdržel se, 0 nepřítomen.'],
]:
    m.append(r)
m.column_dimensions['A'].width = 24; m.column_dimensions['B'].width = 140
m['A1'].font = Font(bold=True, size=13)
k = wb.create_sheet('Kandidáti')
Q = list(PRAVA) + list(SOC)
hdr = ['Obvod', 'Jméno', 'Navrhla', 'Pohlaví (podle jména)', 'Výsledek'] + [f'{q}: {({**PRAVA, **SOC})[q]}' for q in Q] + ['Práva: Pro', 'Práva: Pro+Proti', 'Index práva', 'Sociální: Pro', 'Sociální: Pro+Proti', 'Index sociální']
k.append(hdr)
for c, a, _ in sorted(odp.values(), key=lambda t: (int(t[0]['OBVOD']), t[0]['PRIJMENI'])):
    vys = 'zvolen/a' if zvol(c) else ('prohrál/a ve 2. kole' if fin(c) else 'nepostoupil/a')
    pp = sum(a[q] == 'Pro' for q in PRAVA); pn = sum(a[q] in ('Pro', 'Proti') for q in PRAVA)
    sp = sum(a[q] == 'Pro' for q in SOC); sn = sum(a[q] in ('Pro', 'Proti') for q in SOC)
    k.append([int(c['OBVOD']), f"{c['JMENO']} {c['PRIJMENI']}", cv.get(c['NSTRANA'], c['NSTRANA']), 'žena' if zena(c) else 'muž', vys]
             + [a[q] or '—' for q in Q] + [pp, pn, round(pp / pn, 3) if pn else None, sp, sn, round(sp / sn, 3) if sn else None])
for col in ('C', 'D', 'E'): k.column_dimensions[col].width = 14
k.column_dimensions['B'].width = 30
for i, h in enumerate(hdr, 1):
    k.cell(1, i).font = Font(bold=True); k.cell(1, i).alignment = Alignment(wrap_text=True, vertical='top')
k.freeze_panes = 'C2'
s = wb.create_sheet('Souhrny')
s.append(['Skupina', 'Kandidátů s odpovědí', 'Průměrný index práva', 'Průměrný index sociální'])
def radek(nazev, f):
    x = [a for c, a, _ in odp.values() if f(c)]
    p = [v for v in (idx(a, PRAVA) for a in x) if v is not None]; so = [v for v in (idx(a, SOC) for a in x) if v is not None]
    s.append([nazev, len(x), round(statistics.mean(p), 3) if p else None, round(statistics.mean(so), 3) if so else None])
for nazev, f in [('všichni', lambda c: True), ('zvolení', zvol), ('nezvolení', lambda c: not zvol(c)), ('prohráli ve 2. kole', lambda c: fin(c) and not zvol(c)),
                 ('ženy', zena), ('muži', lambda c: not zena(c))] + [(f'navrhla {st}', (lambda c, st=st: cv.get(c['NSTRANA'], c['NSTRANA']) == st)) for st in ['ANO', 'SPD', 'ODS', 'STAN', 'KDU-ČSL', 'TOP 09']]:
    radek(nazev, f)
s.append([]); s.append(['Otázka', 'Zvolení: Pro / Pro+Proti', 'Prohráli ve 2. kole: Pro / Pro+Proti', 'Všichni: Pro / Pro+Proti'])
for q, t in {**PRAVA, **SOC}.items():
    def pr(f):
        v = [a[q] for c, a, _ in odp.values() if f(c) and a[q] in ('Pro', 'Proti')]; return f"{sum(x == 'Pro' for x in v)} / {len(v)}"
    s.append([f'{q}: {t}', pr(zvol), pr(lambda c: fin(c) and not zvol(c)), pr(lambda c: True)])
s.column_dimensions['A'].width = 50
for col in 'BCD': s.column_dimensions[col].width = 30
ob = wb.create_sheet('Obvody – Istanbul')
ob.append(['Obvod', 'Odcházející senátor/ka', 'Hlas k Istanbulské úmluvě 24. 1. 2024 (senat.cz)', 'Nově zvolený/á', 'Kalkulačka otázka 16 (Istanbul)', 'Obhájil/a'])
for o in range(3, 82, 3):
    p = puvodni[o]; n = novy[o]
    ob.append([o, p, VYKLAD.get(VOTE.get(p, '?'), 'nebyl/a v Senátu'), f"{n['JMENO']} {n['PRIJMENI']}", ans.get(o, {}).get(16) or 'nevyplnil/a',
               'ano' if p.split()[-1] == n['PRIJMENI'].split()[-1] else ''])
for col, w in zip('ABCDEF', (8, 28, 30, 30, 34, 10)): ob.column_dimensions[col].width = w
for ws_ in (s, ob):
    for c_ in ws_[1]: c_.font = Font(bold=True)
wb.save(os.path.join(H, 'kalkulacka-index-detail.xlsx'))
print('\nuloženo: kalkulacka-index-detail.xlsx')

print('\n=== Citlivost: přísnější index „Pro / všechny otázky bloku“ (Nevím i nevyplněno = ne pro) ===')
def idx2(a, qs): return sum(a[q] == 'Pro' for q in qs) / len(qs)
for nazev, f in [('zvolení', zvol), ('nezvolení', lambda c: not zvol(c)), ('prohráli ve 2. kole', lambda c: fin(c) and not zvol(c)), ('ženy', zena), ('muži', lambda c: not zena(c))]:
    x = [a for c, a, _ in odp.values() if f(c)]
    print(f"{nazev:22s} n={len(x):3d} práva {100*statistics.mean(idx2(a, PRAVA) for a in x):5.1f} %  sociální {100*statistics.mean(idx2(a, SOC) for a in x):5.1f} %")

# --- 27 nově zvolených: rozpis odpovědí na 13 otázek (podklad pro rafičky 0–27) -> ../tecky.json ---
OTAZKY = {8: 'Souhlasím se zařazením menstruačních potřeb do snížené sazby DPH.', 16: 'Má Česko ratifikovat Istanbulskou úmluvu?',
          23: 'Má Česko umožnit uzavřít manželství párům stejného pohlaví?', 24: 'Má Česko zakotvit právo na interrupci v Ústavě?',
          2: 'Má stát více pomáhat s financováním výstavby obecních nájemních bytů?', 6: 'Mají lidé s vyššími příjmy platit vyšší daně než dnes?',
          14: 'Lidé pracující pro aplikace na rozvoz jídla či taxislužby by měli mít statut zaměstnanců se všemi sociálními jistotami namísto práce na živnostenský list.',
          15: 'Minimální mzda by měla ze zákona dosahovat alespoň poloviny průměrné mzdy v zemi.', 17: 'Mají se dlouhodobě bezvýsledné exekuce automaticky zastavovat?',
          9: 'Má stát zrušit koncesionářské poplatky za veřejnoprávní média?', 13: 'Má Česko usnadnit zaměstnávání pracovníků ze zemí mimo EU?',
          31: 'Má Česko zastropovat důchodový věk na 65 letech?', 33: 'Má Česko zavést nulovou DPH na léky na předpis?'}
zvol27 = [c for c in reg if zvol(c)]
odp_by = {(c['OBVOD'], c['PRIJMENI']): a for c, a, _ in odp.values()}
out = []
print('\n=== 27 zvolených: pro / proti / nevím / nevyplnili ===')
for q, t in OTAZKY.items():
    cnt = collections.Counter()
    for c in zvol27:
        a = odp_by.get((c['OBVOD'], c['PRIJMENI']))
        v = 'nevyplnil' if a is None else ('pro' if a[q] == 'Pro' else 'proti' if a[q] == 'Proti' else 'nevim')
        cnt[v] += 1
    out.append(dict(q=q, otazka=t, blok='prava' if q in PRAVA else 'socialni' if q in SOC else 'dalsi', pro=cnt['pro'], proti=cnt['proti'], nevim=cnt['nevim'], nevyplnil=cnt['nevyplnil'], celkem=27))
    print(f"{q:2d} pro {cnt['pro']:2d} | proti {cnt['proti']:2d} | nevím {cnt['nevim']:2d} | nevyplnil {cnt['nevyplnil']} | {t[:60]}")
import json as _j
# jednotliví zvolení (veřejné odpovědi z kalkulačky): tečka = senátor/ka
senatori = []
for c in sorted(zvol27, key=lambda c: int(c['OBVOD'])):
    a = odp_by.get((c['OBVOD'], c['PRIJMENI']))
    def kod(q):
        if a is None: return 'x'
        return 'p' if a[q] == 'Pro' else 'n' if a[q] == 'Proti' else 'v'
    senatori.append(dict(jmeno=f"{c['JMENO']} {c['PRIJMENI']}", obvod=int(c['OBVOD']), odpovedi={str(q): kod(q) for q in OTAZKY}))
_j.dump(dict(zdroj='Volební kalkulačka 2026 – Senát (KohoVolit.eu / DataTimes), odpovědi kandidátů; zvolení podle ČSÚ',
             legenda=dict(p='pro', n='proti', v='nevím / nemám jednoznačný názor', x='kalkulačku nevyplnil/a'), otazky=out, senatori=senatori),
        open(os.path.join(H, '..', 'tecky.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
