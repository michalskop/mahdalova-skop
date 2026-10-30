"""Z tabulek ČSÚ „Projekce obyvatelstva v krajích ČR 2024–2080“ (14 souborů 13005224NN.xlsx, list Tab. 1 – počet obyvatel
podle věku k 1. 1., obě pohlaví) spočítá změnu počtu obyvatel a dospělých (18+) mezi 1. 1. 2024 a 1. 1. 2081 → data/kraje_projekce.csv.
Spuštění: python -X utf8 -I kraje_to_csv.py <složka s rozbaleným 13005224_data.zip>"""
import csv, glob, os, sys
import openpyxl

here = os.path.dirname(os.path.abspath(__file__))
src = sys.argv[1]
SHORT = {"Hlavní město Praha": "Praha", "Středočeský kraj": "Středočeský", "Jihočeský kraj": "Jihočeský", "Plzeňský kraj": "Plzeňský",
         "Karlovarský kraj": "Karlovarský", "Ústecký kraj": "Ústecký", "Liberecký kraj": "Liberecký", "Královéhradecký kraj": "Královéhradecký",
         "Pardubický kraj": "Pardubický", "Kraj Vysočina": "Vysočina", "Jihomoravský kraj": "Jihomoravský", "Olomoucký kraj": "Olomoucký",
         "Zlínský kraj": "Zlínský", "Moravskoslezský kraj": "Moravskoslezský"}
out = []
for f in sorted(glob.glob(os.path.join(src, "13005224??.xlsx"))):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    data = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    name = data[0][1].split(":")[0]
    hdr, tot = data[1], data[2]
    cols = [i for i, h in enumerate(hdr) if h is not None and i >= 1]
    i0, i1 = cols[0], cols[-1]
    assert str(hdr[i0]).startswith("2024") and hdr[i1] == 2081, (name, hdr[i0], hdr[i1])
    adults = lambda c: sum(r[c] for r in data[3:] if isinstance(r[0], (int, float)) and r[0] >= 18 and r[c] is not None)
    a, b, ad0, ad1 = tot[i0], tot[i1], adults(i0), adults(i1)
    out.append({"kraj": SHORT[name], "kraj_plny": name, "obyv_2024": round(a), "obyv_2081": round(b),
                "zmena_obyv_pct": round(100 * (b / a - 1), 1), "dospeli_2024": round(ad0), "dospeli_2081": round(ad1),
                "zmena_dospeli_pct": round(100 * (ad1 / ad0 - 1), 1)})
assert len(out) == 14
with open(os.path.join(here, "data", "kraje_projekce.csv"), "w", encoding="utf-8", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(out[0]))
    w.writeheader(); w.writerows(out)
for r in sorted(out, key=lambda r: -r["zmena_obyv_pct"]):
    print(r["kraj"], r["zmena_obyv_pct"], r["zmena_dospeli_pct"])
