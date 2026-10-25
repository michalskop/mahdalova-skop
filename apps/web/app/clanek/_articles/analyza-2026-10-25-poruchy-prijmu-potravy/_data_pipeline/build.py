"""Poruchy příjmu potravy u dětí a dospívajících v Česku – všechna čísla článku.

Vstupy (vstupy/):
- Datovy-souhrn-NR-04-25-poruchy-prijmu-potravy-2026-02.xlsx – ÚZIS / NZIP, datový souhrn
  „Poruchy příjmu potravy“: unikátní pacienti s dg. F50 (hlavní i vedlejší) v péči hrazené
  z veřejného zdravotního pojištění, ambulance i lůžka, 2010–2025, věk 0–18 a 19+.
  https://www.nzip.cz/data/2672-poruchy-prijmu-potravy-datovy-souhrn (staženo 24. 10. 2026)
- Datovy-souhrn-SSS-07-06-vykaz-a013-psychiatrie-2026-01.xlsx – ÚZIS / NZIP, výkaz A013
  ambulantní psychiatrie 2011–2025 (počty za pracoviště, jeden pacient může být započten víckrát).
  https://www.nzip.cz/data/1796-vykaz-a013-psychiatrie-datovy-souhrn (staženo 24. 10. 2026)
- obyvatelstvo_0_19_csu.csv – výřez z ČSÚ 130181 (viz extract_obyvatelstvo.py), stav k 31. 12. 2008–2024.
- hbsc_cz_body_image_dieting.csv – HBSC ČR: Data Browser HBSC, API WHO/Europe a grafy českého
  týmu HBSC (UP Olomouc) na zdravagenerace.cz/report/obezita/ (sestaveno 24. 10. 2026).

Výstup: vystup.txt (všechna čísla v textu), ../data/graf_lecba.json, ../data/graf_hbsc.json
Spuštění: python -I build.py
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = Path(__file__).parent
VST = HERE / "vstupy"
DATA = HERE.parent / "data"

DIVKY = "#f76800"   # brandOrange.6 (ženy)
CHLAPCI = "#06677d"  # brandTeal.7 (muži)
GRID = "#e8e3d2"

out_lines = []


def out(s=""):
    print(s)
    out_lines.append(s)


def cz(x, d=0):
    return f"{x:,.{d}f}".replace(",", " ").replace(".", ",")


# ------------------------------------------------------------ NR-04-25: pacienti s F50
wb = openpyxl.load_workbook(VST / "Datovy-souhrn-NR-04-25-poruchy-prijmu-potravy-2026-02.xlsx", read_only=True, data_only=True)
ws = wb["Data"]
pac = {}  # rok -> dict
for r in ws.iter_rows(values_only=True):
    if isinstance(r[0], int) and 2000 < r[0] < 2100:
        vals = [c for c in r[1:] if c is not None]
        m018, m19, z018, z19, c018, c19 = vals[:6]
        assert m018 + z018 == c018 and m19 + z19 == c19, r
        pac[r[0]] = dict(chlapci=m018, divky=z018, celkem=c018, dospeli=c19)
roky = sorted(pac)
assert roky[0] == 2010 and roky[-1] == 2025

# ------------------------------------------------------------ ČSÚ: dívky a chlapci 0–18
pop = {}
for r in csv.DictReader(open(VST / "obyvatelstvo_0_19_csu.csv", encoding="utf-8")):
    if int(r["vek"]) <= 18:
        k = (int(r["stav_k"][:4]), r["pohlavi"])
        pop[k] = pop.get(k, 0) + int(r["pocet"])


def stredni(rok, pohl):
    """Střední stav ≈ průměr stavů k 31. 12. předchozího a daného roku."""
    a, b = pop.get((rok - 1, pohl)), pop.get((rok, pohl))
    return (a + b) / 2 if a and b else None


out("== Pacienti s poruchou příjmu potravy (F50), 0–18 let, NZIP NR-04-25 ==")
out("rok | dívky | chlapci | celkem | dívky na 10 000 dívek 0–18 | podíl dívek")
na10k = {}
for y in roky:
    p = pac[y]
    s = stredni(y, "dívky")
    if s:
        na10k[y] = p["divky"] / s * 10000
    out(f"{y} | {p['divky']} | {p['chlapci']} | {p['celkem']} | "
        f"{cz(na10k[y], 1) if y in na10k else '–'} | {cz(p['divky'] / p['celkem'] * 100, 1)} %")

d10, d19, d20, d21, d25 = (pac[y]["divky"] for y in (2010, 2019, 2020, 2021, 2025))
out()
out(f"Dívky 2010 → 2025: {d10} → {d25} = {cz(d25 / d10, 2)}× (+{cz((d25 / d10 - 1) * 100)} %)")
out(f"Dívky 2010 → 2019 (před covidem): {d10} → {d19} = +{cz((d19 / d10 - 1) * 100)} %")
out(f"Dívky 2020 → 2021: {d20} → {d21} = +{cz((d21 / d20 - 1) * 100)} % (+{d21 - d20})")
out(f"Dívky 2019 → 2021: {d19} → {d21} = +{cz((d21 / d19 - 1) * 100)} %")
out(f"Dívky 2021 → 2025: {d21} → {d25} = +{cz((d25 / d21 - 1) * 100)} %")
c10, c25 = pac[2010]["chlapci"], pac[2025]["chlapci"]
out(f"Chlapci 2010 → 2025: {c10} → {c25} = {cz(c25 / c10, 2)}×")
out(f"Podíl dívek mezi pacienty 0–18: 2010 {cz(d10 / pac[2010]['celkem'] * 100, 1)} %, 2025 {cz(d25 / pac[2025]['celkem'] * 100, 1)} %")
out(f"Dospělí 19+ 2010 → 2025: {pac[2010]['dospeli']} → {pac[2025]['dospeli']} = +{cz((pac[2025]['dospeli'] / pac[2010]['dospeli'] - 1) * 100)} %")
out(f"Všichni pacienti 2025: {pac[2025]['celkem'] + pac[2025]['dospeli']}")
out()
p10, p24 = stredni(2010, "dívky"), stredni(2024, "dívky")
out(f"Dívky 0–18 v populaci (střední stav): 2010 {cz(p10)}, 2024 {cz(p24)} (+{cz((p24 / p10 - 1) * 100, 1)} %)")
out(f"Dívky v léčbě na 10 000 dívek 0–18: 2010 {cz(na10k[2010], 1)}, 2019 {cz(na10k[2019], 1)}, 2024 {cz(na10k[2024], 1)} "
    f"= {cz(na10k[2024] / na10k[2010], 2)}× (2010–2024; populace za 2025 ČSÚ zatím nezveřejnil)")
out(f"Dívky 2024 v léčbě: {pac[2024]['divky']} = {cz(pac[2024]['divky'] / p24 * 100, 2)} % dívek 0–18 "
    f"(zhruba {cz(pac[2024]['divky'] / p24 * 1000, 1)} z tisíce)")

# ------------------------------------------------------------ A013: ambulance podle věku
wb = openpyxl.load_workbook(VST / "Datovy-souhrn-SSS-07-06-vykaz-a013-psychiatrie-2026-01.xlsx", read_only=True, data_only=True)
ws = wb["tabulky absolutni"]
a013 = {}
for r in ws.iter_rows(values_only=True):
    lab = next((c for c in r if isinstance(c, str)), "")
    if lab.startswith("Neurotické poruchy: poruchy příjmu potravy:"):
        vals = [c for c in r if isinstance(c, (int, float))]
        a013[lab.split(":")[-1].strip()] = dict(zip(range(2011, 2026), [round(v) for v in vals]))
out()
out("== Ambulantní psychiatrie, F50 (výkaz A013, počty za pracoviště) ==")
for k in ("celkem", "0 - 14 let", "15 - 19 let", "nově zjištěná onemocnění"):
    s = a013[k]
    out(f"{k}: 2011 {s[2011]}, 2019 {s[2019]}, 2021 {s[2021]}, 2025 {s[2025]} = {cz(s[2025] / s[2011], 2)}× (2011–2025)")
out("Pozn.: rok 2014 je ve výkazu neúplný (ÚZIS ho v ročence nepublikoval).")

# ------------------------------------------------------------ HBSC
hb = list(csv.DictReader(open(VST / "hbsc_cz_body_image_dieting.csv", encoding="utf-8-sig")))


def hbsc(ukazatel_zacatek, zdroj_zacatek=""):
    res = {}
    for r in hb:
        if r["ukazatel"].startswith(ukazatel_zacatek) and r["zdroj_url"].startswith(zdroj_zacatek):
            res[(int(r["rok"]), int(r["vek"]), r["pohlavi"])] = float(r["hodnota_pct"])
    return res


nejisti = hbsc("normální/nižší váha podle BMI, ale cítí se příliš tlustý")
tlusti = hbsc("cítí se příliš tlustý", "https://data-browser")
dieta = hbsc("drží dietu")
out()
out("== HBSC ČR ==")
out("Normální nebo nižší váha podle BMI, a přesto se cítí příliš tlustá (podíl ze všech dětí dané skupiny), dívky:")
for vek in (11, 13, 15):
    out(f"  {vek} let: " + ", ".join(f"{y} {cz(nejisti[(y, vek, 'dívky')], 1)} %" for y in (2002, 2006, 2010, 2014, 2018, 2022)))
out("  chlapci 15 let: " + ", ".join(f"{y} {cz(nejisti[(y, 15, 'chlapci')], 1)} %" for y in (2002, 2014, 2022)))
out("Cítí se příliš tlustý/á (trochu + hodně), Data Browser HBSC:")
for vek in (11, 13, 15):
    out(f"  {vek} let dívky: " + ", ".join(f"{y} {cz(tlusti[(y, vek, 'dívky')])} %" for y in (2014, 2018, 2022))
        + " | chlapci: " + ", ".join(f"{y} {cz(tlusti[(y, vek, 'chlapci')])} %" for y in (2014, 2018, 2022)))
out("Drží dietu nebo dělá něco pro zhubnutí (2014, API WHO/Europe):")
for vek in (11, 13, 15):
    out(f"  {vek} let: dívky {cz(dieta[(2014, vek, 'dívky')])} %, chlapci {cz(dieta[(2014, vek, 'chlapci')])} %")

(HERE / "vystup.txt").write_text("\n".join(out_lines) + "\n", encoding="utf-8")

# ------------------------------------------------------------ graf 1: léčba
vals = []
for y in roky:
    vals.append({"rok": str(y), "skupina": "dívky", "počet": pac[y]["divky"], "pořadí": 0})
    vals.append({"rok": str(y), "skupina": "chlapci", "počet": pac[y]["chlapci"], "pořadí": 1})
popisky = [{"rok": str(y), "počet": pac[y]["divky"], "text": cz(pac[y]["divky"])} for y in (2010, 2019, 2021, 2025)]
sort = [str(y) for y in roky]
graf_lecba = {
    "$schema": "https://vega.github.io/schema/vega-lite/v5.json",
    "title": {
        "text": "Dívek léčených s poruchou příjmu potravy je {třikrát víc|" + DIVKY + "} než v roce 2010",
        "subtitle": "Děti a dospívající do 18 let s diagnózou poruchy příjmu potravy (F50) v ambulantní nebo nemocniční péči hrazené pojišťovnami. Každé dítě se v roce počítá jednou.",
    },
    "_source": "ÚZIS, NZIP – datový souhrn Poruchy příjmu potravy (NR-04-25, aktualizace 2026); výpočet DataTimes.cz",
    "width": "container",
    "height": 320,
    "layer": [
        {
            "data": {"values": vals},
            "mark": {"type": "bar", "width": {"band": 0.75}},
            "encoding": {
                "x": {"field": "rok", "type": "ordinal", "title": None, "sort": sort,
                      "axis": {"labelAngle": 0, "ticks": False,
                               "labelExpr": "(+datum.value) % 5 == 0 || datum.value == '2021' || datum.value == '2025' ? datum.value : ''"}},
                "y": {"field": "počet", "type": "quantitative", "stack": "zero", "title": None,
                      "axis": {"gridColor": GRID, "format": ",d"}},
                "color": {"field": "skupina", "type": "nominal",
                          "scale": {"domain": ["dívky", "chlapci"], "range": [DIVKY, CHLAPCI]},
                          "legend": {"orient": "top-left", "title": None, "direction": "horizontal"}},
                "order": {"field": "pořadí"},
                "tooltip": [{"field": "rok", "title": "Rok"}, {"field": "skupina", "title": "Skupina"},
                            {"field": "počet", "title": "Pacientů", "format": ",d"}],
            },
        },
        {
            "data": {"values": popisky},
            "mark": {"type": "text", "dy": 9, "baseline": "top", "fontSize": 12, "fontWeight": "bold", "color": "#ffffff"},
            "encoding": {"x": {"field": "rok", "type": "ordinal", "sort": sort},
                         "y": {"field": "počet", "type": "quantitative"},
                         "text": {"field": "text"}},
        },
    ],
    "config": {"view": {"stroke": None}, "locale": {"number": {"decimal": ",", "thousands": " ", "grouping": [3]}}},
}

# ------------------------------------------------------------ graf 2: HBSC
hv = []
for vek, barva in ((11, "11 let"), (13, "13 let"), (15, "15 let")):
    for y in (2002, 2006, 2010, 2014, 2018, 2022):
        hv.append({"rok": str(y), "věk": barva, "podíl": nejisti[(y, vek, "dívky")]})
graf_hbsc = {
    "$schema": "https://vega.github.io/schema/vega-lite/v5.json",
    "title": {
        "text": "Dívek s normální váhou, které se cítí {tlusté|" + DIVKY + "}, je dnes méně než před dvaceti lety",
        "subtitle": "Podíl dívek, které mají podle vlastní výšky a váhy normální nebo nižší váhu, a přesto o sobě říkají, že jsou trochu nebo hodně tlusté. Výzkum zdraví školáků HBSC.",
    },
    "_source": "HBSC ČR, Univerzita Palackého v Olomouci – report Zdravá generace: Obezita; výpočet DataTimes.cz",
    "width": "container",
    "height": 300,
    "data": {"values": hv},
    "mark": {"type": "line", "point": {"size": 60, "filled": True}, "strokeWidth": 3},
    "encoding": {
        "x": {"field": "rok", "type": "ordinal", "title": None, "axis": {"labelAngle": 0, "ticks": False}},
        "y": {"field": "podíl", "type": "quantitative", "title": None, "scale": {"domain": [0, 35]},
              "axis": {"gridColor": GRID, "labelExpr": "datum.value + ' %'", "tickCount": 4}},
        "color": {"field": "věk", "type": "nominal",
                  "scale": {"domain": ["11 let", "13 let", "15 let"], "range": ["#fda668", "#f76800", "#994800"]},
                  "legend": {"orient": "top-right", "title": None, "direction": "horizontal"}},
        "tooltip": [{"field": "rok", "title": "Rok"}, {"field": "věk", "title": "Věk"},
                    {"field": "podíl", "title": "Podíl dívek (%)", "format": ".1f"}],
    },
    "config": {"view": {"stroke": None}, "locale": {"number": {"decimal": ",", "thousands": " ", "grouping": [3]}}},
}

DATA.mkdir(exist_ok=True)
for name, g in (("graf_lecba.json", graf_lecba), ("graf_hbsc.json", graf_hbsc)):
    (DATA / name).write_text(json.dumps(g, ensure_ascii=False, indent=1), encoding="utf-8")
print("grafy zapsány")
